#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
作品4：Python 数据清洗与自动化日报
输入：data/*.csv
输出：output/excel/经营日报_YYYYMMDD.xlsx
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os, sys
from datetime import datetime

BASE = os.path.join(os.path.dirname(__file__), "..")
DATA = f"{BASE}/data"
OUT  = f"{BASE}/output/excel"
os.makedirs(OUT, exist_ok=True)


def load_and_clean():
    """加载并清洗数据"""
    users = pd.read_csv(f"{DATA}/users.csv", parse_dates=["register_time"])
    products = pd.read_csv(f"{DATA}/products.csv")
    orders = pd.read_csv(f"{DATA}/orders.csv", parse_dates=["order_time"])
    items = pd.read_csv(f"{DATA}/order_items.csv")

    # 清洗 1：去除重复
    orders = orders.drop_duplicates(subset=["order_id"])
    items = items.drop_duplicates()

    # 清洗 2：异常值（金额<=0、数量<=0）
    items = items[(items.amount > 0) & (items.qty > 0)]

    # 清洗 3：剔除孤儿明细
    items = items[items.order_id.isin(orders.order_id)]

    # 清洗 4：关联维度（订单城市为准，用户渠道为辅）
    df = items.merge(orders, on="order_id").merge(
        users[["user_id","channel"]], on="user_id", how="left")
    return df


def build_daily_report(df, report_date="2025-12-15"):
    """生成指定日期的经营日报"""
    df["date"] = df.order_time.dt.date
    target = pd.to_datetime(report_date).date()
    today = df[df.date == target]
    yesterday = df[df.date == target - pd.Timedelta(days=1)]

    wb = Workbook()

    # 样式
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    title_font = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    bold = Font(name="微软雅黑", size=10, bold=True)
    normal = Font(name="微软雅黑", size=10)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row, ncols):
        for c in range(1, ncols+1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

    # ---- Sheet1 经营概览 ----
    ws = wb.active
    ws.title = "经营概览"
    ws["A1"] = f"电商经营日报 - {report_date}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:D1")

    kpis = [
        ("指标", "今日", "昨日", "环比"),
        ("GMV(元)", today.amount.sum(), yesterday.amount.sum(),
         f"{(today.amount.sum()-yesterday.amount.sum())/max(yesterday.amount.sum(),1)*100:.1f}%"),
        ("订单数", today.order_id.nunique(), yesterday.order_id.nunique(), "-"),
        ("购买用户数", today.user_id.nunique(), yesterday.user_id.nunique(), "-"),
        ("客单价(元)", round(today.amount.sum()/max(today.order_id.nunique(),1),2), "-", "-"),
    ]
    for r, row in enumerate(kpis, start=3):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            cell.alignment = center
            if r == 3:
                cell.font = header_font; cell.fill = header_fill
            else:
                cell.font = bold if c == 1 else normal

    for col, w in zip("ABCD", [16, 14, 14, 12]):
        ws.column_dimensions[col].width = w

    # ---- Sheet2 品类日报 ----
    ws2 = wb.create_sheet("品类日报")
    cat_today = today.groupby("category").agg(
        gmv=("amount","sum"), orders=("order_id","nunique"), users=("user_id","nunique")
    ).reset_index().sort_values("gmv", ascending=False)
    headers = ["品类", "GMV(元)", "订单数", "购买用户数", "GMV占比"]
    for c, h in enumerate(headers, start=1):
        ws2.cell(row=1, column=c, value=h)
    style_header(ws2, 1, len(headers))
    total = cat_today.gmv.sum()
    for r, row in enumerate(cat_today.itertuples(index=False), start=2):
        ws2.cell(row=r, column=1, value=row.category).font = normal
        ws2.cell(row=r, column=2, value=round(row.gmv,2)).font = normal
        ws2.cell(row=r, column=3, value=row.orders).font = normal
        ws2.cell(row=r, column=4, value=row.users).font = normal
        ws2.cell(row=r, column=5, value=f"{row.gmv/max(total,1)*100:.1f}%").font = normal
    for col, w in zip("ABCDE", [12, 14, 10, 12, 12]):
        ws2.column_dimensions[col].width = w

    # ---- Sheet3 城市日报 ----
    ws3 = wb.create_sheet("城市日报")
    city_today = today.groupby("city").agg(
        gmv=("amount","sum"), orders=("order_id","nunique")
    ).reset_index().sort_values("gmv", ascending=False)
    headers = ["城市", "GMV(元)", "订单数"]
    for c, h in enumerate(headers, start=1):
        ws3.cell(row=1, column=c, value=h)
    style_header(ws3, 1, len(headers))
    for r, row in enumerate(city_today.itertuples(index=False), start=2):
        ws3.cell(row=r, column=1, value=row.city).font = normal
        ws3.cell(row=r, column=2, value=round(row.gmv,2)).font = normal
        ws3.cell(row=r, column=3, value=row.orders).font = normal
    for col, w in zip("ABC", [12, 14, 10]):
        ws3.column_dimensions[col].width = w

    out = f"{OUT}/经营日报_{report_date.replace('-','')}.xlsx"
    wb.save(out)
    print(f"日报已生成: {out}")
    return out


if __name__ == "__main__":
    df = load_and_clean()
    # 取数据中最后一个有订单的日期作为示例
    last_date = df.order_time.max().date()
    build_daily_report(df, str(last_date))
